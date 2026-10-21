# Lecture d'un classeur Excel déposé pour géolocaliser les sondages d'une affaire (import Excel de
# l'édition sur la carte, US-CARTE-26 du front). Rien n'est enregistré : le classeur est rendu
# feuille par feuille, cellules en texte, pour que le front propose l'association des colonnes.

import math
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.errors import InvalidRequest

MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
# Une feuille de mesures (CPT) dépasse le millier de lignes : au-delà, elle est tronquée et signalée.
MAX_ROWS = 5000
MAX_COLUMNS = 100
OPEN_XML_EXTENSIONS = (".xlsx", ".xlsm")
LEGACY_EXTENSION = ".xls"


@dataclass(frozen=True)
class Sheet:
    name: str
    rows: list[list[str]]
    truncated: bool


# Texte d'une cellule : nombres en notation décimale avec un point (jamais « 1e-05 », que le front
# refuse), dates au format français, vide pour une cellule sans valeur.
def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "VRAI" if value else "FAUX"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(Decimal(repr(value)).normalize(), "f") if math.isfinite(value) else ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y") if value.time() == time() else value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value)


def _without_trailing_blanks(cells: list[str]) -> list[str]:
    while cells and cells[-1] == "":
        cells.pop()
    return cells


def _sheet(name: str, rows_of_values) -> Sheet:
    rows: list[list[str]] = []
    truncated = False
    for values in rows_of_values:
        if len(rows) == MAX_ROWS:
            truncated = True
            break
        rows.append(_without_trailing_blanks([cell_text(value) for value in values[:MAX_COLUMNS]]))
    while rows and not rows[-1]:
        rows.pop()
    return Sheet(name, rows, truncated)


def _read_open_xml(content: bytes) -> list[Sheet]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        raise InvalidRequest("Le fichier n'est pas un classeur Excel lisible.") from error

    try:
        return [_sheet(worksheet.title, worksheet.iter_rows(values_only=True)) for worksheet in workbook.worksheets]
    finally:
        workbook.close()


# Valeur Python d'une cellule xlrd : le format .xls range les dates parmi les nombres (le type de
# cellule les distingue) et les valeurs logiques en 0/1 ; une cellule en erreur (#DIV/0!…) est vide.
def _legacy_value(cell: xlrd.sheet.Cell, datemode: int) -> object:
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        except xlrd.xldate.XLDateError:
            return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    return cell.value


def _legacy_rows(worksheet: xlrd.sheet.Sheet, datemode: int):
    for index in range(worksheet.nrows):
        yield [_legacy_value(cell, datemode) for cell in worksheet.row(index)]


def _read_legacy(content: bytes) -> list[Sheet]:
    # xlrd lève des exceptions de types variés (XLRDError, CompDocError, struct.error, IndexError…)
    # sur un fichier abîmé ou qui n'est pas un classeur : toutes reviennent à un fichier illisible.
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as error:
        raise InvalidRequest("Le fichier n'est pas un classeur Excel lisible.") from error

    try:
        return [
            _sheet(worksheet.name, _legacy_rows(worksheet, workbook.datemode))
            for worksheet in (workbook.sheet_by_index(index) for index in range(workbook.nsheets))
        ]
    except Exception as error:
        raise InvalidRequest("Le fichier n'est pas un classeur Excel lisible.") from error
    finally:
        workbook.release_resources()


def read_workbook(file_name: str, content: bytes) -> list[Sheet]:
    name = file_name.lower()
    if not name.endswith((*OPEN_XML_EXTENSIONS, LEGACY_EXTENSION)):
        raise InvalidRequest("Seuls les classeurs Excel .xlsx, .xlsm et .xls sont lus.")
    if not content:
        raise InvalidRequest("Le fichier est vide.")
    if len(content) > MAX_WORKBOOK_BYTES:
        raise InvalidRequest("Le fichier dépasse 10 Mo.")

    return _read_legacy(content) if name.endswith(LEGACY_EXTENSION) else _read_open_xml(content)
